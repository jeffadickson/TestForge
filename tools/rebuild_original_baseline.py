from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path
import json

OUT=Path(__file__).resolve().parents[1]/'build'/'baseline-v0.1'; OUT.mkdir(parents=True,exist_ok=True)
doc=Document()
sec=doc.sections[0]; sec.page_width=Inches(8.5); sec.page_height=Inches(11)
sec.top_margin=sec.bottom_margin=Inches(.7); sec.left_margin=sec.right_margin=Inches(.8)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
    s=doc.styles[name]; s.font.name='Calibri'; s.font.color.rgb=RGBColor(0,0,0)
for el in doc.styles.element.xpath('.//w:pBdr'):
    el.getparent().remove(el)
doc.styles['Normal'].font.size=Pt(11)
doc.styles['Normal'].paragraph_format.space_after=Pt(6)
doc.styles['Normal'].paragraph_format.line_spacing=1.04
doc.styles['Title'].font.size=Pt(27)
doc.styles['Heading 1'].font.size=Pt(18)
doc.styles['Heading 2'].font.size=Pt(12)
doc.styles['Heading 2'].paragraph_format.space_before=Pt(10)
footer=sec.footer.paragraphs[0]; footer.alignment=2
r=footer.add_run('Flight Hardware Test Framework  |  0.1  |  '); r.font.size=Pt(9)
fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); footer._p.append(fld)
reqs=[]
def p(t,style=None): return doc.add_paragraph(t,style)
def h(t): doc.add_heading(t,2)
def page(t):
    para=doc.add_heading(t,1); para.paragraph_format.page_break_before=True
def table(headers,rows,widths):
    t=doc.add_table(rows=1, cols=len(headers)); t.autofit=False
    for i,(s,w) in enumerate(zip(headers,widths)):
        t.columns[i].width=Inches(w); t.rows[0].cells[i].text=s
    for row in rows:
        cells=t.add_row().cells
        for i,v in enumerate(row): cells[i].text=str(v)
    for ri,row in enumerate(t.rows):
        for ci,c in enumerate(row.cells):
            c.width=Inches(widths[ci]); c.vertical_alignment=1
            pr=c._tc.get_or_add_tcPr(); borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'4'); e.set(qn('w:color'),'D9D9D9'); borders.append(e)
            pr.append(borders)
            mar=OxmlElement('w:tcMar')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge); e.set(qn('w:w'),'90'); e.set(qn('w:type'),'dxa'); mar.append(e)
            pr.append(mar)
            shade=OxmlElement('w:shd'); shade.set(qn('w:fill'),'E6E6E6' if ri==0 else 'FFFFFF'); pr.append(shade)
            for para in c.paragraphs:
                para.paragraph_format.space_after=Pt(2)
                for run in para.runs: run.font.size=Pt(10); run.bold=ri==0
    repeat=OxmlElement('w:tblHeader'); t.rows[0]._tr.get_or_add_trPr().append(repeat)
    p('')
    return t

doc.add_heading('Flight Hardware Test Framework',0)
p('Research and requirements baseline', 'Subtitle')
p('Version 0.1  •  14 October 2026  •  Prepared for Jeff Dickson')
h('Recommended direction')
p('Build a controlled procedure and evidence system around a durable local station runtime. Use a web browser for authoring, execution, witnessing, review and analysis. The station must execute approved work and preserve its record without a connection to the central database, identity service or Internet. Synchronization and reconciliation are explicit workflow activities.')
p('Evaluate OpenTAP and OpenHTF as replaceable execution adapters before committing to a new sequencer. Use Epsilon3 as the strongest procedure-management benchmark. Keep the authoritative execution journal independent of a general application framework, including Loom or Frappe.')
h('Scope and status')
p('This baseline specifies human and automated flight-hardware testing from procedure preparation through execution, reconciliation, analysis and archival closure. It covers manual steps, Python, VISA and USB instruments, calibration evidence, redlines, separate QA and operator signatures, multiple approvers, read-only observers and long-running offline work.')
p('The requirements below are proposed system requirements, not claims of an implemented or flight-qualified system. Numerical targets are initial engineering acceptance targets to be ratified against representative test campaigns. Product findings are documentation-based; no vendor product was bench-qualified for this study. Institutional test, safety, quality and records policies govern the eventual deployment.')
h('Principal design decisions')
for t in [
'Persist locally before acknowledging a step, signature or result. Replicate centrally afterward, even when the network is healthy.',
'Separate execution state, measurement verdict, synchronization state and review disposition. A passed measurement does not mean an accepted test record.',
'Give one station authority over each active run and its hardware. Disconnected copies cannot independently command the same article.',
'Never automatically replay an ambiguous hardware operation after a crash. Recover evidence, establish hardware state and obtain an explicit disposition.',
'Bind signatures to exact versions and evidence. Preserve corrections and retests as new records rather than overwriting history.']: p(t,'List Bullet')
h('Document map')
p('Product research and framework assessment; station architecture and synchronization; lifecycle and data contracts; 96 requirements with acceptance criteria; verification scenarios; staged delivery and unresolved decisions; source register.')

page('Research findings and product fit')
p('The market separates instrument-oriented test executives from human procedure and quality systems. Combining those capabilities is feasible, but an execution engine does not establish the complete approval and evidence model. “Not established” below means the reviewed documentation does not prove the requested behavior; it is not a claim that the product cannot provide it.')
table(['Candidate','Documented capability','Fit for this framework'],[
('NI TestStand','Sequences, Python and other code adapters, result processing and database reporting. [S01–S03]','Instrument automation benchmark; controlled web procedures and approval reconciliation need separate evaluation.'),
('Epsilon3','Controlled procedures, redlines, approvals, APIs and offline procedure features. [S04–S08]','Closest functional benchmark for the human workflow; proprietary platform.'),
('OpenTAP','Extensible .NET engine, Python integration, resource and result plugins. [S09–S12]','Leading C# execution candidate; quality workflow remains application work.'),
('OpenHTF','Python phases, measurements, hardware plugs and attachments. Apache 2.0. [S13]','Leading Python execution candidate; add durable procedure and approval layer.'),
('OpenC3 COSMOS','Browser Script Runner, Python/Ruby, suites, telemetry and script lifecycle. [S14–S16]','Strong spacecraft command and telemetry candidate; evaluate license and edition boundaries.'),
('Yamcs and Open MCT','Command, stack and script activities; telemetry and mission visualization. [S17–S19]','Integration targets when spacecraft command and telemetry dominate; not a proven complete QA procedure system.')
],[1.05,2.8,3.05])
h('How to interpret the evidence')
p('Documentation establishes that the relevant building blocks exist. It does not establish power-loss durability, offline identity guarantees, safe recovery of irreversible actions, or approval behavior after conflicting edits. Require witnessed fault-injection demonstrations and record the exact product edition, version and plugin set.')
p('The absence of an Internet dependency must be tested from a cold restart. A page left open before disconnecting is insufficient evidence of a deployable offline system.')

page('Commercial systems and procedure management')
h('NI TestStand')
p('TestStand is the NI product most likely intended by “Test Runner.” It supports reusable code modules, including Python, and configurable reporting and database logging. On-the-fly result processing can write during execution. Its Offline Results Processing Utility converts raw results into reports and database tables later or on another machine. This is a relevant precedent for decoupling execution from central storage. [S01–S03]')
p('Do not equate deferred result processing with the complete requested offline workflow. Demonstrate offline identity, operator and QA signatures, approved redlines, insertions, calibration snapshots, run ownership and post-reconnect reconciliation. Also test the exact persistence boundary: collected in memory, written to a file and durably committed are different states. A custom operator interface is not evidence of browser-native operation.')
h('Epsilon3')
p('Epsilon3 is the closest documented match for controlled human procedures: released revisions, review, sign-offs, redlines and external integrations. Its published changes explicitly describe completing steps offline, starting prepared procedures offline, adding comments and attachments offline, and synchronizing afterward. It also advertises on-premise deployment. [S04–S08]')
p('Those features substantially overlap the requested system; the rationale for building is control over behavior, source and integration, rather than a claim that no product exists. Ask for a demonstration of a fully disconnected restart, Python execution at the stand, local multi-user witnesses, expired credentials, stale calibration, conflicting redlines and a partially uploaded waveform. Older changelogs establish feature history, not present limits or crash guarantees.')
h('Evaluation script for both vendors')
p('Use the same representative procedure: identify an engineering article, verify calibration, acquire an operator and independent QA sign-off, run a bounded voltage measurement, insert an approved diagnostic step, lose the network, restart the station, finish locally, reconnect, resolve a stale reference, upload a large attachment and separately approve analysis.')
p('Demand a complete machine-readable export containing identities, timestamps, versions, revisions, attempts, original and corrected values, signatures and raw attachments. A PDF report alone is not sufficient portability. Request edition-specific licensing, disconnected activation requirements, deployment topology, data export rights and a documented recovery procedure; pricing was not established by this review.')

page('Open execution and spacecraft platforms')
h('OpenTAP')
p('The core is .NET-based under MPL 2.0. Its plans contain steps and resources; result listeners separate acquisition from reporting. Official documentation describes SQLite and PostgreSQL listeners and a Python plugin. This is a strong match for a C# host with Python test code. [S09–S11]')
p('The Runner documentation exposes programmatic sessions and run events, but also references registration and licensing. Do not assume every editor, driver, runner or cloud component shares the core license. Prefer a small owned adapter to the core until dependencies and deployment rights are enumerated. A result listener is useful integration plumbing; the station must still prove commit-before-acknowledgment and recovery behavior. [S12]')
h('OpenHTF')
p('OpenHTF is a Python library with phase functions, measurement validation, instrument plugs and binary attachments. Its Apache 2.0 repository provides inspectable code and examples. It is attractive when test engineers want conventional Python. Its run verdict and teardown machinery are useful, but a persisted multi-day human workflow, independent sign-offs and offline reconciliation remain application requirements. [S13]')
p('Evaluate one measurement sequence through OpenHTF and one through OpenTAP using the same event contract. Judge cancellation, hangs, hardware resource exclusion, result streaming and dependency packaging. Do not select solely from syntax or an attractive editor.')
h('OpenC3 COSMOS')
p('COSMOS provides a browser Script Runner, Python/Ruby scripting, suites and script lifecycle controls. Its line stepping and command/telemetry integration make it relevant for spacecraft integration and test. Human-facing execution controls need policy constraints in a controlled flight test. [S14–S15]')
p('Current license documentation describes a Builder License with hosted-service and commercial redistribution restrictions, plus commercial and evaluation alternatives. Enterprise features include user accounts and RBAC. Treat this as source-visible software with specific restrictions, not as equivalent to OpenHTF’s Apache license or OpenTAP’s MPL core. Inspect the exact chosen release and plugins before adoption. [S16]')
h('Yamcs and Open MCT')
p('Yamcs tracks command, stack and script activities. Open MCT is a mission-data visualization framework; AMMOS identifies Yamcs and Open MCT as core elements of its multi-center mission-control direction. They are relevant integrations for telemetry and commanding, rather than a substitute for the required quality record. Requirements for calibrated bench instruments and independent procedure closure remain ours to verify. [S17–S19]')

page('Loom and Frappe assessment')
h('Loom source inspection')
p('The inspected artifact is Loom_Framework_v0.12.zip, current version 12. README, architecture and requirements documents, DocumentService.cs, Storage.cs, SqliteStore.cs, Program.cs and data-type documentation were examined. This was source inspection, not a new execution or security qualification. [S26]')
p('Useful existing pieces include C# metadata and typed repositories, SQLite transactions with audit writes, optimistic version checks, stable field identities, schema migration support, binary compression and external data references. SqliteStore explicitly enables WAL; the application should additionally set and verify its intended durability settings instead of relying on provider defaults.')
p('The current host uses simulated roles and loopback restrictions, explicitly not production login. Its document lifecycle is Draft, Submitted and Cancelled. Durable jobs, a transactional outbox, generalized workflows, synchronization and production deployment are listed as open work. Queries are bounded in-process scans; the store refuses lists above 10,000 records. Combined inline binary data has an 8 MiB document limit. These are concrete reasons not to use the existing generic record store as a high-volume acquisition journal or waveform archive.')
h('Frappe assessment')
p('Frappe offers a web application framework, document metadata, permissions and workflow transitions. Its database API defines transaction behavior, and background jobs use a queue with after-commit enqueue support. These capabilities suit central authoring, asset registers, assigned review work and reporting. [S20–S22]')
p('A queued job is not automatically a recoverable hardware operation. A Frappe site on a laptop would still need its own server dependencies and a designed synchronization protocol. The reviewed docs do not establish a turnkey disconnected flight-test journal with split-brain protection. Model multiple required approvers as explicit approval records and policy checks; do not infer them from a single workflow-state label.')
table(['Layer','Proposed starting point','Boundary'],[
('Station runtime','.NET service and dedicated SQLite journal','May reuse selected Loom types; must own recovery, identity, resource fencing and sync.'),
('Test scripts','Versioned Python workers','OpenHTF or OpenTAP adapter; instruments remain behind controlled ownership.'),
('Central application','Frappe pilot or dedicated API','Own authoring and review; choose one authority for approvals and records.'),
('Loom integration','Versioned application adapter','Exchange metadata and result references through an API; no direct uncontrolled database writes.')
],[1.25,2.4,3.25])
p('Recommendation: keep the station contract framework-neutral. Frappe is the shorter route to a central administrative application; Loom is a plausible longer-term C# reuse path. Neither choice removes the need to engineer and verify the local execution core.')

page('Deployment and local authority')
h('Three operating modes')
p('Connected station: the browser talks to the local station service; the station commits locally and replicates centrally. Isolated station: the same browser and service run with prepared packages and local identities. Isolated site: multiple browser clients on a local network use one site station authority while the enterprise network is unavailable. A browser on another device cannot access that station when all communication to it is lost.')
h('Component responsibilities')
table(['Component','Responsibility'],[
('Browser application','Procedure display, inputs, live results, role actions and synchronization status; locally served assets without CDN dependencies.'),
('Station authority','Validate step guards; journal events; manage signatures and authority; supervise Python workers; hold instrument ownership.'),
('Local storage','Transactional event journal and outbox; content-addressed raw files; prepared packages, reference snapshots and authorization bundles.'),
('Instrument gateway','Resource locks and command authorization; protocol adapters; operation IDs and reconciliation hooks. Hardware protection remains independently enforced.'),
('Synchronization service','Resumable replication, hash verification, deduplication, evidence completeness and receipts; no hardware command replay.'),
('Central services','Released procedures, users, assets, requirement links, review tasks, search, analysis records and retention. Derived views can be rebuilt from retained events.')
],[1.55,5.35])
h('Why the station service is necessary')
p('The browser is the user interface, not the sole durable recorder or universal instrument driver. Browser-managed storage has quota and eviction behavior; persistent storage reduces some eviction risk but does not eliminate user deletion or device loss. Native VISA and arbitrary USB drivers belong on the station. Browser closure must not silently kill an acquisition. [S23, S24]')
p('Use a local database transaction to commit the accepted action, corresponding event and outbox entry together. Store large artifacts separately with staged writes, verified hashes and explicit incomplete states. SQLite WAL with synchronous FULL is a candidate; durability still depends on the storage stack honoring flushes. Keep the database on a local supported filesystem and use coordinated backup rather than copying only a live database file. [S25]')
h('Ownership and availability')
p('One station owns a run and its resources. Offline authority is an explicit bounded grant, not a short central lease that expires and lets a second station seize the hardware. Handoff requires release or physical isolation of the previous commander. A central outage is tolerated; loss of the local evidence store or required hardware protection is a different failure and must inhibit new work.')

page('Disconnected preparation and synchronization')
h('Before disconnecting')
p('Prepare a signed package containing the released procedure and subprocedures, step identities, required documents, scripts and locked dependencies, driver versions, limits, configuration, calibration snapshots, approved overrides, authorized identities and roles, expiry policy and expected storage budget. Verify all files and execute a readiness check with the network disabled. A preparation receipt records exactly what is available.')
h('During execution')
p('Each event carries an event ID, run ID, station identity, ownership epoch, monotonically increasing local sequence, event schema version, actor, role, local timestamp, clock-quality information, causation ID, payload digest and previous-event digest. Integrity chains make alterations detectable against retained checkpoints; they do not make an administrator-controlled machine intrinsically tamper-proof.')
p('The station acknowledges an action only after the necessary record is durable. A worker records command intent before issuing a hardware operation and records the observed outcome afterward. If the computer fails between these records, the outcome is unknown. A database transaction cannot make a physical side effect atomic.')
h('On reconnecting')
for t in [
'Authenticate the station and exchange protocol versions, ownership and last verified receipts. Do not change the active package.',
'Send immutable events and artifact chunks with stable IDs. The server accepts retries idempotently and flags the same ID with different content.',
'Verify sequence continuity, signatures, package hashes and all required artifacts. Receipt of metadata alone is not evidence completeness.',
'Compare cached authorization, calibration and release state with current records. Create reconciliation tasks for changes; preserve the original execution context.',
'Issue a durable reconciliation receipt. Retain local evidence until central completeness and backup/retention conditions are met.',
'Make central closure available only when reconciliation and required reviews are complete. Upload failure never initiates a new measurement.']: p(t,'List Number')
h('Conflict policy')
p('Independent comments and attachments may merge by identity. Conflicting signed results, redlines, approvals or ownership claims must not use last-writer-wins. Preserve both branches, quarantine the disputed transition and assign a named resolver. A newly released procedure must not silently replace the version used by an offline run.')
h('Offline authorization tradeoff')
p('Immediate central revocation and indefinite disconnected authorization cannot both be guaranteed. Authorize a bounded offline period and permitted actions in advance. Preserve signatures made within the local authorization context; flag subsequently discovered revocations for review. Emergency safing must remain available when normal credentials or authorization expire.')

page('Lifecycle and evidence model')
h('Independent state dimensions')
table(['Dimension','Representative states','Meaning'],[
('Procedure','Draft, in review, released, superseded, withdrawn','Defines what was approved to execute.'),
('Execution','Prepared, ready, running, held, recovery required, completed, aborted','Describes activity and hardware-control progress.'),
('Verdict','Not evaluated, pass, fail, inconclusive, invalid','Describes measurement or requirement evaluation.'),
('Replication','Local only, queued, uploading, uploaded, reconciled, conflict','Describes evidence availability and consistency.'),
('Review','Open, execution reviewed, analysis open, analysis approved, closed, reopened','Describes acceptance and outstanding obligations.')
],[1.05,3.5,2.35])
p('Execution completion requires all executed, skipped or deferred steps to have an explicit disposition and the final hardware configuration to be recorded. Execution review can be performed locally when policy allows. Final central closure additionally requires complete reconciliation, analysis acceptance and resolution or authorized disposition of deviations. Closed records reopen through amendments; historical approvals remain visible.')
h('Canonical entities')
p('Procedure and ProcedureRevision; StepDefinition with permanent identity; Requirement and VerificationLink; ExecutionPackage; TestArticle and ConfigurationSnapshot; Station and AuthorityGrant; Run and StepAttempt; CommandIntent and CommandOutcome; Measurement and ArtifactManifest; CalibrationSnapshot; IdentityGrant and Signature; Redline and Deviation; ReviewTask and AnalysisRevision; SyncReceipt and ReconciliationIssue; ClosureManifest.')
p('A run points to one execution package. Each repeated step creates a new attempt linked to its parent step and preceding attempts. Each measurement points to an attempt, instrument identity, calibration context, units and limits revision. Each analysis points to immutable input hashes and a code/environment version. A signature points to the exact subject revision and evidence digest it approves, not merely a mutable record ID.')
h('API and worker contract')
p('Expose versioned operations for preparing packages, starting runs, proposing actions, signing subjects, submitting measurements, reserving artifacts, appending events, uploading chunks and acknowledging receipts. Mutation requests carry a request ID, expected version, actor authority and run ownership. Read APIs expose pending and durable states distinctly.')
p('Python workers receive a validated input document, package digest, attempt ID and approved resource handles. They emit structured measurements, progress, artifacts and terminal outcomes. stdout is diagnostic output, not the authoritative measurement protocol. Database adapters use outboxes and stable IDs; a remote database outage never forces an instrument operation to run again.')

page('Requirements conventions and allocation')
p('The following 96 requirements form the proposed first baseline. Each states an observable obligation and its acceptance check. They are grouped by ownership and mapped to end-to-end qualification scenarios later in the document. All require implementation evidence before release; none is declared satisfied by this study.')
table(['Priority','Interpretation'],[
('P0','Required before any controlled flight-hardware pilot. Development may use simulators while these are unfinished.'),
('P1','Required for the complete operational framework; staged after the core pilot if the pilot scope explicitly excludes it.')
],[.8,6.1])
table(['Group','Allocated owner','Primary evidence'],[
('PROC','Procedure service','Versioning, import and guard tests'),('EXEC','Station runtime','State and interruption tests'),('AUTH','Identity and approvals','Role, signature and revocation tests'),('CHG','Change control','Redline, retest and invalidation tests'),('HW','Instrument gateway','Bench and fault-injection tests'),('CAL','Asset and calibration adapters','Reference validity and recall tests'),('DATA','Evidence service','Integrity, precision and restore tests'),('OFF','Local runtime','Disconnected cold-start and restart tests'),('SYNC','Replication service','Partition, duplicate and conflict tests'),('CLOSE','Review and analysis','Closure-blocking and traceability tests'),('SEC','Platform security','Access, isolation and tampering tests'),('OPS','Operations and user interface','Load, installation and recovery demonstrations')
],[.8,2.1,4.0])
h('Verification method')
p('Acceptance checks are minimum evidence, not complete test procedures. Use automated unit and integration tests for guards, versions and deduplication; process and power interruption for persistence; real representative equipment for command behavior; and witnessed operator/QA demonstrations for human workflows. Preserve the software build, environment, test inputs, expected outcome, actual outcome and reviewer for every acceptance result.')

def group(title, prefix, owner, items):
    page(title)
    p('Allocation: '+owner+'.')
    for i,(label,priority,shall,accept) in enumerate(items,1):
        rid=f'{prefix}-{i:02d}'; h(f'{rid}  {label}  {priority}')
        if prefix=='OPS': doc.paragraphs[-1].paragraph_format.space_before=Pt(7)
        p('The system shall '+shall)
        para=p('Acceptance: '+accept); para.paragraph_format.space_after=Pt(7)
        reqs.append(dict(id=rid,title=label,priority=priority,owner=owner,requirement='The system shall '+shall,acceptance=accept))

group('Controlled procedures','PROC','Procedure service',[
('Stable identity','P0','assign permanent IDs to procedures, revisions and steps, independent of displayed numbering.','Insert a step and verify historical signatures and result links still reference the original steps.'),
('Immutable release','P0','freeze a released revision and its dependent content; edits shall create a new revision.','Modify a draft and verify an existing released package and its digest remain unchanged.'),
('Structured steps','P0','support instructions, typed inputs, measurements, scripts, hold points, witnesses, attachments and subprocedures.','Execute a representative procedure containing every supported step type with required inputs enforced.'),
('Execution guards','P0','enforce prerequisites, allowed branches, loops and mandatory hold points in the runtime.','Attempt a direct API bypass of a hold point and confirm rejection and an audit event.'),
('Release validation','P0','reject missing references, invalid units, unresolved required roles and unsupported package dependencies.','Run invalid-package fixtures and verify release is blocked with actionable diagnostics.'),
('Configuration applicability','P0','restrict procedure use by article identity, configuration and approved applicability rules.','Start against a mismatched hardware revision and confirm an explicit blocking result.'),
('Import and export','P1','import structured procedures as drafts and export full content with stable identities and revision history.','Round-trip a procedure and identify any unsupported imported content before release.'),
('Requirement mapping','P0','link steps and analyses to versioned verification requirements and methods.','Change a requirement revision and verify affected released procedures and evidence are flagged for review.')])

group('Execution and recovery','EXEC','Station runtime',[
('Run identity','P0','assign unique run and attempt IDs with article, station, package and personnel context.','Start two offline runs and verify their imported records cannot collide.'),
('Durable transitions','P0','commit an accepted transition and its event before returning success to the browser.','Terminate immediately after acknowledgment and recover the same accepted state.'),
('Explicit outcomes','P0','distinguish pass, fail, inconclusive, error, invalid, aborted, skipped and not evaluated outcomes.','Inject each outcome and verify it cannot become pass through missing data or an exception.'),
('Hold and resume','P0','apply declared hold and browser-disconnect policies, and recheck prerequisites before resumption.','Close the browser during acquisition, then change an instrument during a hold; verify the declared response and blocked resume.'),
('Unknown outcome','P0','enter recovery required after an interrupted operation whose hardware effect cannot be established.','Crash after command transmission but before result persistence; verify no automatic retransmission.'),
('Retry classification','P0','require each operation to declare read-only, idempotent or nonrepeatable behavior with an approved retry policy.','Verify retries are bounded and a nonrepeatable operation cannot use generic automatic retries.'),
('Resource ownership','P0','permit one controlling authority per run and exclusive resource set, including during disconnection.','Attempt a second controller and confirm hardware access is denied until a valid handoff.'),
('Handover record','P0','record shift handover, active configuration, held operations and transfer of operator responsibility.','Resume after a personnel change and verify the new operator acknowledges outstanding conditions.')])

group('Identity signatures and collaboration','AUTH','Identity and approval service',[
('Individual identity','P0','authenticate individual operators, witnesses and approvers without shared sign-off accounts.','Verify two people using one workstation retain distinct authenticated signature records.'),
('Independent witness','P0','enforce separate operator and QA identities when a step requires independence.','Give one account both roles and confirm it still cannot satisfy both signatures.'),
('Approval policies','P0','support named approvers, role slots, all-of and explicitly configured quorum rules with ordering.','Exercise three required approvers, rejection and optional observers without premature approval.'),
('Signature meaning','P0','bind each signature to identity, role, meaning, timestamp, policy and exact subject/evidence digest.','Alter the signed content and verify the old signature cannot approve the new revision.'),
('Reauthentication','P0','require a configured confirmation and reauthentication policy for consequential signatures.','Expire a signing session and verify a normal UI click cannot silently sign.'),
('Offline authority','P0','validate locally provisioned, bounded identity and role grants for disconnected actions.','Cold-start offline and verify permitted signing works while an expired or out-of-scope grant fails.'),
('Observers and comments','P0','separate viewing, commenting, executing, witnessing and approving permissions.','Verify a read-only collaborator neither mutates the run nor becomes a required approver.'),
('Delegation and revocation','P0','record authorized delegation and reconcile newly discovered revocations without erasing prior signatures.','Revoke an offline user centrally and verify reconnect creates an attributable review issue.')])

group('Redlines deviations and retests','CHG','Change control service',[
('Proposed redlines','P0','record redlines as explicit proposals with original content, proposed content, reason and affected scope.','Show the exact before/after change and proposer for an inserted or altered step.'),
('Approval before use','P0','block affected execution until the required redline approvals are complete.','Propose a changed limit while running and verify the original approved limit remains operative.'),
('Run and master scope','P0','distinguish run-only changes from proposed changes to the master procedure.','Approve a run insertion and verify other runs and the released master do not change.'),
('Evidence preservation','P0','retain original observations and append corrections with reason, author and time.','Correct a signed numeric result and verify both values and affected approvals remain traceable.'),
('Skip and defer','P0','require reason, authority and disposition for skipped, deferred or not-applicable required steps.','Attempt closure with an unauthorized skip and confirm it is blocked.'),
('Retest linkage','P0','create a new attempt for every retest and retain prior failed and invalid attempts.','Pass a retest and verify the failed attempt remains in reports and requirement disposition.'),
('Impact invalidation','P0','mark affected downstream evidence and approvals for reevaluation after consequential changes.','Change a calibration correction or limit and verify dependent analysis approval becomes stale.'),
('Nonconformance','P0','link deviations and nonconformances to article, run, affected steps and authorized dispositions.','Accept a documented waiver and verify the measured failure is preserved rather than relabeled pass.')])

group('Python and instrument control','HW','Workers and instrument gateway',[
('Script provenance','P0','execute only approved script packages with pinned dependencies and recorded hashes in flight mode.','Change one script byte or dependency and confirm execution is refused.'),
('Process isolation','P0','run scripts in supervised worker processes with bounded resources and separate credentials.','Hang or crash a worker and verify the UI, journal and supervision remain available.'),
('Driver abstraction','P0','provide versioned adapters for VISA, serial, USB and required vendor APIs with simulated equivalents.','Run the same measurement contract against a simulator and representative approved equipment.'),
('Instrument identity','P0','verify actual device identity and configured resource mapping before permitting commands.','Swap two USB devices or change an address and verify the wrong instrument is blocked.'),
('Operation journal','P0','persist command intent, operation ID and observed outcome without claiming atomic hardware execution.','Interrupt at each dispatch boundary and verify an honest completed or unknown classification.'),
('Bounded cancellation','P0','specify command timeouts, cancellation escalation and equipment-specific recovery actions.','Block a driver call and demonstrate the approved timeout and independent recovery behavior.'),
('Safety boundary','P0','respect independent interlocks and enforce approved command envelopes and operating modes.','Request an out-of-envelope command and verify rejection; prove a failed process cannot defeat the hardware interlock.'),
('Data protocol','P0','ingest typed worker measurements and artifacts through a validated protocol separate from diagnostic logs.','Send malformed units, duplicate samples and a truncated artifact; verify explicit rejection or incomplete status.')])

group('Calibration and reference information','CAL','Asset and reference adapters',[
('Asset chain','P0','record instrument, channel, probe, cable, fixture and transducer identities needed for measurement traceability.','Reconstruct the complete measurement chain from the saved run without live asset-system access.'),
('Calibration snapshot','P0','retain certificate revision, validity interval, correction data, uncertainty and source retrieval metadata.','Disconnect the calibration database and reproduce the context from the prepared package.'),
('Validity checks','P0','check calibration at preparation and required use points, including expiry during a long test.','Cross an expiry boundary and verify the approved hold or disposition policy is enforced.'),
('Reference freshness','P0','apply configured freshness and expiry policies to each cached external reference.','Use a stale required reference offline and verify the action is blocked or explicitly authorized.'),
('Controlled exceptions','P0','require scoped approval for calibration or reference exceptions and record affected measurements.','Verify an exception for one instrument and run does not authorize another.'),
('Recall impact','P0','identify past measurements and analyses affected by a later calibration recall or out-of-tolerance finding.','Recall a certificate and generate the exact affected evidence set with review tasks.'),
('Measurement quality','P0','retain range, resolution, uncertainty model and approved guard-band decision rules when applicable.','Evaluate values near a limit and verify the approved rule and its revision are reported.'),
('External adapters','P1','provide read adapters with provenance, timeouts, schema validation and independently testable mappings.','Change an upstream schema and verify failure is explicit without modifying prior snapshots.')])

group('Scientific data and evidence','DATA','Evidence service',[
('Measurement representation','P0','preserve raw values, units, precision, validity flags and evaluated limits without silent conversion.','Round-trip large integers, decimals, timestamps and invalid readings without loss or false pass.'),
('Raw and derived data','P0','retain raw acquisitions separately from reductions, corrections and analysis outputs.','Recompute a derived value from retained raw evidence and identify the exact algorithm revision.'),
('Large artifacts','P0','stream waveforms, images and binary files to bounded local storage with size and digest manifests.','Acquire a 10 GB artifact without loading it entirely into memory and detect a truncated final chunk.'),
('Timing evidence','P0','record source time, receipt time, clock source, synchronization quality and sequence order.','Move the wall clock backward and verify event order remains reconstructable and uncertainty is visible.'),
('Evidence integrity','P0','verify content digests and signed checkpoints when storing, synchronizing and exporting evidence.','Alter or remove a retained event or artifact and verify the integrity check detects it.'),
('Storage exhaustion','P0','monitor capacity and apply a validated reserve policy before new operations can exhaust evidence storage.','Fill the disk during acquisition and verify explicit incomplete evidence and the approved safe response.'),
('Portable record','P0','export human-readable reports and machine-readable events, schemas, manifests and raw artifacts.','Open an export on an isolated clean system and verify its completeness without a running server.'),
('Retention and restoration','P0','retain protected records and verify coordinated backup and restore of journals and artifact files.','Restore a backup and compare all committed evidence and manifest hashes with the source.')])

group('Disconnected operation','OFF','Local station runtime',[
('Local web application','P0','serve all required UI assets and APIs locally without runtime Internet dependencies.','Start the browser and station after a reboot with every external network disconnected.'),
('Prepared package','P0','verify package completeness, signatures, dependencies, reference validity and storage before offline readiness.','Remove each required package element in turn and verify readiness fails.'),
('Local source of record','P0','record every acknowledged input, signature, transition and result locally before replication.','Disable the central database throughout a run and recover every acknowledged record locally.'),
('Crash durability','P0','recover committed events and identify incomplete operations after process or power interruption.','Inject interruptions across commit boundaries on the target storage and verify no acknowledged event is lost.'),
('Disconnected duration','P0','support a proposed seven-day isolated campaign within an approved authority and capacity budget.','Complete a seven-day soak without external authentication, licensing, package or time-service dependencies.'),
('Offline collaboration','P0','support authorized local operator and witness use through one station authority, with explicit connectivity status.','Use two browsers on an isolated LAN and reject conflicting execution actions.'),
('Readiness boundaries','P0','distinguish central service failure from unavailable local storage, invalid authority or missing safety prerequisites.','Verify central outage permits prepared work while loss of local durable storage inhibits new operations.'),
('Portable transport','P1','support encrypted, signed export/import bundles when network synchronization is unavailable.','Import the same removable-media bundle twice and verify deduplication and custody records.')])

group('Synchronization and reconciliation','SYNC','Replication service',[
('Transactional outbox','P0','atomically store each accepted event and its replication obligation in the local transaction.','Crash after commit and before upload; verify the event remains queued after restart.'),
('Idempotent ingest','P0','deduplicate retransmissions by stable identity and reject identical IDs with different content.','Retry a batch after losing its acknowledgment and verify exactly one central record per event.'),
('Resumable transfer','P0','resume verified artifact chunks without marking incomplete evidence as complete.','Disconnect repeatedly during a 10 GB upload and verify the final digest and missing-chunk accounting.'),
('Completeness receipt','P0','issue reconciliation receipts only after required sequences, signatures and artifacts validate.','Upload all metadata but omit one attachment and verify closure remains blocked.'),
('Conflict quarantine','P0','preserve conflicting signed records and require explicit resolution instead of last-writer-wins.','Submit two incompatible offline branches and verify both survive with a blocked disputed transition.'),
('Ownership fencing','P0','reject stale ownership epochs and prevent resynchronization from authorizing a second hardware commander.','Reconnect an old station after handoff and verify it cannot resume instrument control.'),
('Reference reconciliation','P0','compare execution snapshots with later releases, recalls and revocations and generate review tasks.','Reconnect after a calibration recall and verify affected runs are flagged without history edits.'),
('Retention acknowledgment','P0','retain local evidence until complete central receipt and configured backup/retention gates are met.','Lose the final receipt or fail the central backup gate and verify local cleanup is refused.')])

group('Review analysis and closure','CLOSE','Review and analysis services',[
('Execution checklist','P0','require step dispositions, required signatures, evidence completeness and final article configuration for execution review.','Leave each item unresolved in turn and verify the review gate explains the blocker.'),
('Separate analysis','P0','track analysis tasks, owners, inputs and approvals independently from execution completion.','Complete a run with analysis pending and verify it is not labeled fully closed.'),
('Reproducible analysis','P0','record analysis code, dependencies, parameters, input hashes and output revisions.','Re-run the approved analysis on retained inputs and compare outputs under stated tolerances.'),
('Approval invalidation','P0','invalidate applicability of analysis approval when its inputs, code or acceptance rules change.','Revise a reduction script and verify a new analysis approval is required.'),
('Closure gate','P0','permit final closure only after reconciliation, required approvals and authorized deviation dispositions.','Attempt closure with one missing approver, unresolved conflict or incomplete upload and confirm denial.'),
('Requirement disposition','P0','report verified, failed, waived, invalid and not-yet-verified requirements with evidence links.','Verify a waived failure remains distinct from a measured pass in the verification matrix.'),
('Amendment lifecycle','P0','reopen closed records through controlled amendments while preserving the prior closure manifest.','Correct an archived analysis and reconstruct both the original and amended acceptance states.'),
('Work queues','P1','provide ownership, due dates, escalation and observer access for execution and analysis action items.','Assign overdue tasks and verify notifications do not silently approve or close them.')])

group('Security and software assurance','SEC','Platform and assurance',[
('Enforced access','P0','enforce role, project and article scope at the API and data-access layers.','Attempt cross-project reads and unauthorized mutations directly through APIs and verify denial.'),
('Local attack surface','P0','authenticate local clients and protect against cross-origin requests, host spoofing and unauthorized LAN access.','Attempt browser-based requests from an untrusted origin and verify commands cannot be issued.'),
('Protected secrets','P0','encrypt sensitive records at rest and in transit, and protect keys and credentials with approved storage and rotation.','Rotate keys without losing historical evidence; verify encrypted storage and transport and absence of secrets in exports.'),
('Approved software','P0','verify signed software packages and prohibit unapproved upgrades during active runs.','Offer a new release during a run and verify it cannot change the pinned execution environment.'),
('Privileged activity','P0','record administrator actions and restrict modification or deletion of authoritative evidence.','Attempt an administrative edit and verify prevention or a detectable controlled amendment.'),
('Worker containment','P0','apply least-privilege filesystem, network and device access to test workers.','Use a test worker to attempt unauthorized file, network and device access and verify rejection.'),
('Assurance record','P0','maintain requirements-to-test traceability, dependency inventory, known limitations and reviewed release evidence.','Audit a candidate release and resolve every P0 requirement to recorded acceptance evidence.'),
('Environment separation','P0','distinguish simulation, development, engineering and flight operation in access rules and stored records.','Attempt to present simulated measurements as flight evidence and verify the provenance is unmistakable.')])

group('Operations usability and performance','OPS','Platform operations and user interface',[
('Operational display','P0','show article identity, revision, active step, ownership, hardware status, local persistence and sync backlog.','An operator identifies the authoritative run and unsynchronized evidence without navigating away.'),
('Clear actions','P0','distinguish hold, resume, abort, skip, repeat and emergency safing with appropriate role and confirmation rules.','Run a witnessed usability exercise and verify no ambiguous button initiates a repeat or closes a failed step.'),
('Accessible interaction','P1','support keyboard operation, readable contrast, non-color status cues and documented supported browsers.','Complete a procedure without a mouse and verify errors remain identifiable without color.'),
('Reference load','P0','meet proposed targets of 10,000 steps per procedure and 100 structured events per second per station.','Run the reference load with a proposed p95 local action acknowledgment below 500 ms, excluding instrument latency.'),
('Startup and recovery','P0','restore the UI and a truthful recovered run state within a proposed 60 seconds on approved hardware.','Restart a loaded station; verify readiness timing without automatically resuming hardware actions.'),
('Observability','P0','expose correlated health, logs and metrics without treating diagnostic telemetry as the evidence record.','Disable telemetry export and verify execution evidence remains complete and locally inspectable.'),
('Install and upgrade','P0','support reproducible offline installation, controlled schema migration and rollback or forward-recovery procedures.','Install on a clean isolated host and rehearse a failed upgrade without losing historical records.'),
('Capacity and support','P0','document storage sizing, backups, restore, instrument compatibility and incident recovery for operators.','A second engineer follows the runbook to restore a station and identify an unresolved command outcome.')])

page('Fault handling and acceptance campaigns')
p('The following campaigns exercise the requirements together. These are proposed tests, not results. Start with simulated hardware and then repeat equipment-specific cases on approved engineering hardware. Capture station build, package hash, fault timing and observed physical state.')
table(['Case','Stimulus and required result','Coverage'],[
('V01','Cold-start without WAN, DNS or central identity. Prepared work and local sign-offs function; invalid grants are blocked.','OFF, AUTH, PROC'),
('V02','Fail the central database for a full run. All acknowledged steps persist locally; no hardware retries arise from upload failure.','OFF, SYNC, EXEC'),
('V03','Kill a worker before dispatch, after dispatch and before outcome commit. Unknown physical effects require reconciliation.','EXEC, HW, DATA'),
('V04','Cut station power at persistence boundaries. No acknowledged event disappears; unfinished artifacts remain incomplete.','OFF, DATA'),
('V05','Lose an upload acknowledgment and resend. One central event is retained; differing payload under the same ID is rejected.','SYNC'),
('V06','Create divergent offline ownership or edits. Both records survive; no second controller or automatic merge is accepted.','EXEC, SYNC, CHG'),
('V07','Expire a calibration during a soak and later recall its certificate. Enforce the local policy and identify affected history.','CAL, CLOSE'),
('V08','Revoke an offline approver centrally. Reconnect creates a review task and preserves the original signing context.','AUTH, SYNC'),
('V09','Fill the disk and corrupt an artifact. New operations are inhibited as specified; evidence cannot be silently completed.','DATA, OFF'),
('V10','Hang a VISA or USB call. Demonstrate timeout, worker containment and equipment-specific independent safing.','HW, EXEC'),
('V11','Redline a limit, insert a step and retest a failure. Preserve original evidence and require dependent reapproval.','PROC, CHG, CLOSE'),
('V12','Finish testing with missing analysis or attachment. Execution can end; final closure remains blocked.','CLOSE, SYNC')
],[.55,5.35,1.0])

page('Qualification scope and delivery sequence')
h('Additional acceptance campaigns')
p('V13 exercises three mandatory approvers, one read-only observer, a rejected approval and a prohibited self-witness. V14 moves the station clock backward and restores a backup; event order and uncertainty must remain explicit. V15 runs the seven-day offline soak and the load targets while uploading large artifacts. V16 exports a closed record and reproduces its analysis on a clean isolated machine. V17 tests API bypass, cross-origin attacks, tampered packages and worker privilege boundaries. V18 rehearses failed installation, migration and rollback. Together with V01–V12, these cover all requirement groups.')
h('Delivery gates')
table(['Gate','Scope','Exit evidence'],[
('A  Architecture spike','One browser, local journal, simulated Python measurement and a server ingest endpoint.','Prove V02, V03 and V05 before expanding UI features.'),
('B  Controlled procedure prototype','Versioned packages, manual steps, operator/QA identities, redlines and local recovery.','Demonstrate full offline execution and approval invalidation using simulated hardware.'),
('C  Engineering bench pilot','Representative VISA/USB devices, calibration adapter, bounded cancellation, raw data and analysis.','Pass relevant equipment fault tests and document safe operating envelopes.'),
('D  Controlled flight pilot','Qualified installation and support, all P0 evidence, reviewed security and project procedures.','Independent project acceptance; no unresolved P0 items within pilot scope.'),
('E  Operational expansion','Additional sites, importers, removable-media exchange, richer work queues and reporting.','Complete P1 acceptance and validate additional deployment profiles.')
],[1.25,2.9,2.75])
h('Prototype that will resolve the largest risks')
p('Use an engineering power-supply and DMM procedure: identify the article; verify instrument identities and calibration; acquire operator and QA permission; command a bounded setpoint; collect readings; approve a diagnostic insertion; disconnect the network; interrupt after dispatch; reconcile actual hardware state; complete locally; synchronize a waveform; generate analysis; obtain independent analysis approval; close the record.')
p('The critical success is correct behavior through interruption and review. A long list of instrument drivers or a sophisticated editor does not compensate for ambiguous command recovery, missing evidence or invalid signatures. Choose OpenTAP versus OpenHTF after both implement this same scenario behind the same station contract.')

page('Open decisions and engineering risks')
table(['Decision','Proposed default','Resolution evidence'],[
('Station platform','Windows first for the actual instrument estate; define Linux support separately.','Inventory VISA stacks, USB drivers and installation rights.'),
('Offline window','Seven days as an initial target, bounded per project.','Longest expected campaign and acceptable revocation/calibration staleness.'),
('Identity offline','Provisioned individual identities and signed bounded grants.','Institutional identity and QA acceptance of offline signatures.'),
('Run ownership','One local authority per run, with explicit handoff.','Procedure for failed-station isolation and controller replacement.'),
('Central framework','Frappe evaluation, with an independent station protocol.','Representative approval, record retention and integration prototype.'),
('Loom reuse','Typed metadata and adapters where useful; dedicated event/blob storage.','Authentication, workflow and storage qualification before expanding responsibility.'),
('Data rates','100 events/s and 10 GB artifact as initial acceptance workload.','Actual waveform rates, retention, channel count and storage sizing.'),
('Safety and stop policy','Equipment-specific; independent hardware protection.','Hazard analysis and tested response for loss of process, power or communication.'),
('Final closure','Central reconciliation required; local execution review allowed by policy.','QA rules for remote-site provisional acceptance and later closure.'),
('Record retention','No deletion of unsynchronized evidence; institutional retention thereafter.','Records schedule, backup destination and recovery objectives.')
],[1.25,2.85,2.8])
h('Failure modes that deserve early design review')
p('A command may have executed despite a timeout. An instrument can retain an output after the controlling process dies. A generic “power off” cleanup may itself be unsafe. A late approval can refer to an obsolete revision. A cached certificate can become invalid while offline. A browser can show a value that never reached durable storage. A centrally accepted event may still lack its raw artifact. A database administrator can rewrite an unauthenticated audit trail. Each has an explicit requirement and acceptance campaign above.')
h('Framework boundary')
p('A generic durable workflow engine may coordinate long-lived human review, but cannot supply physical exactly-once execution. Temporal documents activity retry and idempotency concepts; apply that distinction to any engine selected. Test execution must own the ambiguity between persisted command intent and observable physical effect. [S27]')
p('NASA verification guidance supports a plan-to-evidence-to-disposition lifecycle and traceability to requirements. It informs this structure; applicability of any NASA or institutional requirement must be assigned by the project rather than inferred from product selection. [S28]')

sources=[
('S01','NI TestStand overview','https://www.ni.com/en/shop/electronic-test-instrumentation/application-software-for-electronic-test-and-instrumentation-category/what-is-teststand.html','Product scope and supported code reuse.'),
('S02','NI on-the-fly reporting and logging','https://knowledge.ni.com/KnowledgeArticleDetails?id=kA00Z000000P9LdSAK&l=en-US','Incremental result processing behavior.'),
('S03','NI Offline Results Processing Utility','https://www.ni.com/docs/en-US/bundle/teststand/page/teststand-offline-results-processing-utility.html','Deferred reports and database processing; not a full offline approval guarantee.'),
('S04','Epsilon3 procedures and work instructions','https://www.epsilon3.io/procedures-work-instructions','Controlled revisions, review and redlines.'),
('S05','Epsilon3 offline step execution','https://www.epsilon3.io/behind-the-console/epsilon3-changelog14-offline-step-execution-reviews','Historical evidence for offline step inputs and synchronization.'),
('S06','Epsilon3 offline procedure start','https://www.epsilon3.io/behind-the-console/epsilon3-changelog17-offline-start-email-notifications-more-sso','Historical evidence for prepared offline starts.'),
('S07','Epsilon3 offline attachments','https://www.epsilon3.io/behind-the-console/epsilon3-changelog15-offline-comments-attachments-fancy-headers','Historical evidence for offline comments and attachments.'),
('S08','Epsilon3 integrations and deployment','https://www.epsilon3.io/api-integrations','External data and command integration; deployment options also at https://www.epsilon3.io/pricing .'),
('S09','OpenTAP core repository','https://github.com/opentap/opentap','Core architecture and MPL 2.0 license.'),
('S10','OpenTAP overview','https://doc.opentap.io/User%20Guide/Introduction/Readme.html','Plans, verdicts, resources and result listeners.'),
('S11','OpenTAP Python integration','https://doc.opentap.io/OpenTap.Python/','Python plugin capability.'),
('S12','OpenTAP Runner','https://doc.opentap.io/runner/','Session/API facilities and registration/licensing references; selected component terms require verification.'),
('S13','OpenHTF repository and README','https://github.com/google/openhtf','Apache 2.0; Python phases, measurements, plugs and attachments.'),
('S14','OpenC3 Script Runner','https://docs.openc3.com/docs/tools/script-runner','Browser scripting, execution controls, suites and lifecycle.'),
('S15','OpenC3 scripting API','https://docs.openc3.com/docs/guides/scripting-api','Python and Ruby command and telemetry integration.'),
('S16','OpenC3 licenses','https://docs.openc3.com/docs/meta/licenses','Current license alternatives, restrictions and enterprise feature boundary.'),
('S17','Yamcs activities','https://docs.yamcs.org/yamcs-server-manual/activities/','Command, stack and script activities and status.'),
('S18','NASA Open MCT','https://nasa.github.io/openmct/','Mission visualization framework.'),
('S19','AMMOS mission control products','https://ammos.nasa.gov/mcs/','Yamcs and Open MCT mission-control integration context.'),
('S20','Frappe workflows','https://docs.frappe.io/erpnext/workflows','Document states, transitions and approval workflow configuration.'),
('S21','Frappe database API','https://docs.frappe.io/framework/user/en/api/database','Transaction behavior and database operations.'),
('S22','Frappe background jobs','https://docs.frappe.io/framework/user/en/api/background_jobs','Queue and after-commit execution facilities.'),
('S23','MDN browser storage quotas and eviction','https://developer.mozilla.org/en-US/docs/Web/API/Storage_API/Storage_quotas_and_eviction_criteria','Browser-managed storage persistence and limits.'),
('S24','PyVISA Python backend','https://github.com/pyvisa/pyvisa-py','Message-based instrument transports; per-device support still needs validation.'),
('S25','SQLite write-ahead logging','https://www.sqlite.org/wal.html','WAL durability and operational constraints; also https://www.sqlite.org/pragma.html#pragma_synchronous .'),
('S26','Loom v0.12 local source inspection','Loom_Framework_v0.12.zip  |  version 12','README.md; docs/ARCHITECTURE.md, REQUIREMENTS.md and DATA_TYPES.md; src/Loom.Core/DocumentService.cs and Storage.cs; src/Loom.Sqlite/SqliteStore.cs; src/Loom.Demo/Program.cs.'),
('S27','Temporal activity definitions','https://github.com/temporalio/documentation/blob/main/docs/encyclopedia/activities/activity-definition.mdx','Idempotency and activity execution behavior; architectural reference, not a selected dependency.'),
('S28','NASA product verification guidance','https://www.nasa.gov/reference/5-3-product-verification/','Verification planning, evidence and discrepancy disposition; project applicability is separate.')]
for start in range(0,len(sources),10):
    page('Source register' if start==0 else 'Source register continued')
    if start==0: p('Reviewed 14 October 2026. Vendor and project documentation are primary sources. Historical changelogs establish specific feature precedents; product editions and operational guarantees require demonstration. Architecture, requirements, numerical targets and build recommendations are engineering proposals from this study.')
    for sid,title,url,note in sources[start:start+10]:
        h(sid+'  '+title)
        para=p(''); para.paragraph_format.space_after=Pt(2)
        if url.startswith('https:'):
            from urllib.parse import urlparse
            from docx.opc.constants import RELATIONSHIP_TYPE as RT
            link=OxmlElement('w:hyperlink'); link.set(qn('r:id'),para.part.relate_to(url,RT.HYPERLINK,is_external=True))
            run=OxmlElement('w:r'); props=OxmlElement('w:rPr'); size=OxmlElement('w:sz'); size.set(qn('w:val'),'19'); props.append(size); run.append(props)
            text=OxmlElement('w:t'); text.text='Read source at '+urlparse(url).netloc; run.append(text); link.append(run); para._p.append(link)
        else: para.add_run(url)
        for run in para.runs: run.font.size=Pt(9)
        para=p(note); para.paragraph_format.space_after=Pt(5)
        for run in para.runs: run.font.size=Pt(10)

doc.core_properties.title='Flight Hardware Test Framework Research and Requirements'
doc.core_properties.subject='Offline browser based flight hardware test execution and evidence control'
doc.core_properties.author=''
doc.save(OUT/'Flight_Hardware_Test_Framework_v0.1.docx')
(OUT/'requirements.json').write_text(json.dumps(reqs,indent=2))
print(f'{len(reqs)} requirements; {len(sources)} source entries')
